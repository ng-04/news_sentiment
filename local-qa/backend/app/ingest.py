"""Parse PDF/DOCX/XLSX files into located text units, then split units into chunks."""
import datetime as dt
import io
import posixpath
import re
import zipfile
from dataclasses import dataclass

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from openpyxl import load_workbook
from pypdf import PdfReader
from pypdf.errors import PdfReadError

SUPPORTED_EXTENSIONS = (".pdf", ".docx", ".xlsx", ".xlsm")

# DOCX and XLSX have no fixed pages; these count as one page for the per-session page limit.
DOCX_CHARS_PER_PAGE = 3000
XLSX_ROWS_PER_PAGE = 50


@dataclass
class Unit:
    """A span of source text with a citable location: a PDF page, a DOCX section, or one
    spreadsheet row (sheet + 1-based row number)."""
    text: str
    page: int | None = None
    section: str | None = None
    sheet: str | None = None
    row: int | None = None


@dataclass
class ParsedFile:
    name: str
    units: list[Unit]
    pages: int | None  # real page count for PDFs, None otherwise
    page_equivalent: int  # used for the per-session page limit
    folder: str = ""  # path relative to the chosen folder, "" for the top level


@dataclass
class Chunk:
    file_id: str
    file_name: str
    text: str
    page: int | None
    section: str | None
    chunk_index: int
    folder: str = ""
    sheet: str | None = None
    row_start: int | None = None
    row_end: int | None = None


class ParseError(Exception):
    pass


def parse_file(name: str, data: bytes) -> ParsedFile:
    lower = name.lower()
    if lower.endswith(".pdf"):
        return _parse_pdf(name, data)
    if lower.endswith(".docx"):
        return _parse_docx(name, data)
    if lower.endswith((".xlsx", ".xlsm")):
        return _parse_xlsx(name, data)
    raise ParseError("unsupported file type (only .pdf, .docx, .xlsx and .xlsm)")


def expand_zip(zip_name: str, data: bytes, max_entry_bytes: int, max_entries: int,
               max_total_bytes: int) -> tuple[list[tuple[str, str, bytes]], list[tuple[str, str]]]:
    """Unpack supported files from a zip (e.g. OneDrive's "Download" of a folder).

    Returns ([(file_name, folder, bytes)], [(path, reason) skipped]). Folder paths are kept;
    if the entries don't already sit under one top-level folder, the zip's own name becomes
    that folder, so citations still start with a recognisable folder name. Sizes are checked
    from the zip headers before anything is extracted, then enforced while reading, which
    guards against zip bombs.
    """
    try:
        zf = zipfile.ZipFile(io.BytesIO(data))
    except zipfile.BadZipFile:
        raise ParseError("not a valid .zip file")
    wanted, skipped = [], []
    for info in zf.infolist():
        path = info.filename.replace("\\", "/")
        base = path.rsplit("/", 1)[-1]
        if info.is_dir() or not base or path.startswith("__MACOSX/") or base.startswith((".", "~$")):
            continue
        if not base.lower().endswith(SUPPORTED_EXTENSIONS):
            skipped.append((path, "unsupported file type"))
        elif info.file_size > max_entry_bytes:
            skipped.append((path, f"larger than {max_entry_bytes // (1024 * 1024)} MB"))
        else:
            wanted.append(info)
    if len(wanted) > max_entries:
        skipped += [(i.filename, "file limit reached") for i in wanted[max_entries:]]
        wanted = wanted[:max_entries]
    if sum(i.file_size for i in wanted) > max_total_bytes:
        raise ParseError(f"zip expands to more than {max_total_bytes // (1024 * 1024)} MB")

    tops = {i.filename.replace("\\", "/").split("/", 1)[0] for i in wanted if "/" in i.filename.replace("\\", "/")}
    all_nested = all("/" in i.filename.replace("\\", "/") for i in wanted)
    prefix = "" if (len(tops) == 1 and all_nested) else zip_name.rsplit(".", 1)[0]

    out = []
    for info in wanted:
        path = info.filename.replace("\\", "/")
        folder, _, base = path.rpartition("/")
        with zf.open(info) as fh:
            content = fh.read(max_entry_bytes + 1)
        if len(content) > max_entry_bytes:  # header lied about the size
            skipped.append((path, "larger than allowed"))
            continue
        out.append((base, normalize_folder(f"{prefix}/{folder}" if prefix else folder), content))
    return out, skipped


def normalize_folder(path: str | None) -> str:
    """Clean a client-supplied relative folder path: forward slashes, no leading slash,
    no '.' or '..' segments, at most 300 characters. Returns "" for the top level."""
    if not path:
        return ""
    parts = [p.strip() for p in path.replace("\\", "/").split("/")]
    parts = [p for p in parts if p and p not in (".", "..")]
    return posixpath.join(*parts)[:300] if parts else ""


def _parse_pdf(name: str, data: bytes) -> ParsedFile:
    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise ParseError("PDF is password-protected")
        units = []
        for i, page in enumerate(reader.pages, start=1):
            text = _clean(page.extract_text() or "")
            if text:
                units.append(Unit(text=text, page=i))
        n_pages = len(reader.pages)
    except PdfReadError as e:
        raise ParseError(f"could not read PDF ({e})")
    if not units:
        raise ParseError("no extractable text (the PDF may be scanned images)")
    return ParsedFile(name=name, units=units, pages=n_pages, page_equivalent=n_pages)


def _parse_docx(name: str, data: bytes) -> ParsedFile:
    try:
        doc = Document(io.BytesIO(data))
    except Exception as e:  # python-docx raises several unrelated types for corrupt files
        raise ParseError(f"could not read DOCX ({type(e).__name__})")

    # Walk paragraphs and tables in document order, grouping text under the nearest heading.
    units: list[Unit] = []
    section, buf = None, []

    def flush():
        text = _clean("\n\n".join(buf))
        if text:
            units.append(Unit(text=text, section=section))
        buf.clear()

    for block in doc.element.body.iterchildren():
        tag = block.tag.rsplit("}", 1)[-1]
        if tag == "p":
            para = Paragraph(block, doc)
            style = (para.style.name if para.style is not None else "") or ""
            if style.lower().startswith(("heading", "title")) and para.text.strip():
                flush()
                section = para.text.strip()[:120]
            elif para.text.strip():
                buf.append(para.text)
        elif tag == "tbl":
            table = Table(block, doc)
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells]
                if any(cells):
                    buf.append(" | ".join(cells))
    flush()
    if not units:
        raise ParseError("document contains no text")
    chars = sum(len(u.text) for u in units)
    return ParsedFile(name=name, units=units, pages=None,
                      page_equivalent=max(1, -(-chars // DOCX_CHARS_PER_PAGE)))


def _parse_xlsx(name: str, data: bytes) -> ParsedFile:
    """One unit per data row. The first non-empty row of each sheet is its header, and every
    row is written as "Header: value; ..." so a chunk makes sense on its own."""
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except (zipfile.BadZipFile, KeyError, ValueError, OSError) as e:
        raise ParseError(f"could not read spreadsheet ({type(e).__name__})")
    units: list[Unit] = []
    try:
        for ws in wb.worksheets:
            header = None
            for row_num, row in enumerate(ws.iter_rows(values_only=True), start=1):
                cells = [_cell(v) for v in row]
                if not any(cells):
                    continue
                if header is None:
                    header = [c or _column_letter(i) for i, c in enumerate(cells)]
                    continue
                pairs = [f"{header[i] if i < len(header) else _column_letter(i)}: {c}"
                         for i, c in enumerate(cells) if c]
                units.append(Unit(text="; ".join(pairs), sheet=ws.title, row=row_num))
    finally:
        wb.close()
    if not units:
        raise ParseError("spreadsheet has no data rows")
    return ParsedFile(name=name, units=units, pages=None,
                      page_equivalent=max(1, -(-len(units) // XLSX_ROWS_PER_PAGE)))


def _cell(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, dt.datetime):
        return value.date().isoformat() if value.time() == dt.time() else value.isoformat(sep=" ")
    if isinstance(value, (dt.date, dt.time)):
        return value.isoformat()
    return _clean(str(value)).replace("\n", " ")


def _column_letter(index: int) -> str:
    letters = ""
    index += 1
    while index:
        index, rem = divmod(index - 1, 26)
        letters = chr(65 + rem) + letters
    return f"Column {letters}"


def _clean(text: str) -> str:
    text = text.replace("\x00", "")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


# ---------------------------------------------------------------- chunking

_SEPARATORS = ["\n\n", "\n", ". ", " "]


def chunk_file(file_id: str, parsed: ParsedFile, strategy: str, size: int, overlap: int) -> list[Chunk]:
    if parsed.units and parsed.units[0].sheet is not None:
        return _chunk_rows(file_id, parsed, size)
    chunks: list[Chunk] = []
    for unit in parsed.units:
        if strategy == "by_page":
            pieces = [unit.text] if len(unit.text) <= size * 2 else _fixed(unit.text, size * 2, overlap)
        elif strategy == "by_paragraph":
            pieces = _by_paragraph(unit.text, size, overlap)
        elif strategy == "fixed":
            pieces = _fixed(unit.text, size, overlap)
        else:
            pieces = _recursive(unit.text, size, overlap)
        for piece in pieces:
            piece = piece.strip()
            if piece:
                chunks.append(Chunk(file_id=file_id, file_name=parsed.name, text=piece,
                                    page=unit.page, section=unit.section, chunk_index=len(chunks),
                                    folder=parsed.folder))
    return chunks


def _chunk_rows(file_id: str, parsed: ParsedFile, size: int) -> list[Chunk]:
    """Pack consecutive rows of one sheet into chunks of up to `size` characters. Rows are
    never split, so every chunk can cite an exact sheet and row range. (Each row already
    carries its headers, so overlap isn't needed.)"""
    chunks: list[Chunk] = []
    group: list[Unit] = []

    def flush():
        if group:
            chunks.append(Chunk(file_id=file_id, file_name=parsed.name,
                                text=f"Sheet {group[0].sheet}:\n" + "\n".join(u.text for u in group),
                                page=None, section=None, chunk_index=len(chunks), folder=parsed.folder,
                                sheet=group[0].sheet, row_start=group[0].row, row_end=group[-1].row))
            group.clear()

    length = 0
    for unit in parsed.units:
        if group and (unit.sheet != group[0].sheet or length + len(unit.text) > size):
            flush()
            length = 0
        group.append(unit)
        length += len(unit.text) + 1
    flush()
    return chunks


def _fixed(text: str, size: int, overlap: int) -> list[str]:
    step = max(1, size - overlap)
    return [text[i:i + size] for i in range(0, max(len(text) - overlap, 1), step)]


def _by_paragraph(text: str, size: int, overlap: int) -> list[str]:
    out = []
    for para in re.split(r"\n\s*\n", text):
        out.extend([para] if len(para) <= size else _recursive(para, size, overlap))
    return out


def _recursive(text: str, size: int, overlap: int) -> list[str]:
    """Split on the coarsest separator that works, then pack pieces up to `size` with overlap."""
    pieces = _split(text, size, 0)
    chunks, current = [], ""
    for piece in pieces:
        if current and len(current) + len(piece) > size:
            chunks.append(current)
            tail = current[-overlap:] if overlap else ""
            # Start the overlap at a word boundary so chunks don't begin mid-word.
            if tail and " " in tail:
                tail = tail[tail.index(" ") + 1:]
            current = tail + piece
        else:
            current += piece
    if current.strip():
        chunks.append(current)
    return chunks


def _split(text: str, size: int, level: int) -> list[str]:
    if len(text) <= size:
        return [text]
    if level >= len(_SEPARATORS):
        return [text[i:i + size] for i in range(0, len(text), size)]
    sep = _SEPARATORS[level]
    parts = text.split(sep)
    out = []
    for i, part in enumerate(parts):
        piece = part + (sep if i < len(parts) - 1 else "")
        out.extend([piece] if len(piece) <= size else _split(piece, size, level + 1))
    return out
