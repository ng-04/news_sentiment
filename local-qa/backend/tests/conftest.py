import hashlib
import io
import re

import numpy as np
import pytest
from docx import Document
from openpyxl import Workbook
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

from app.config import Settings

PASSCODE = "open-sesame"


class FakeEmbedder:
    """Deterministic hashed bag-of-words vectors: shared words => higher cosine similarity."""

    dim = 512

    def _vec(self, text: str) -> np.ndarray:
        v = np.zeros(self.dim)
        for word in re.findall(r"[a-z0-9]+", text.lower()):
            if len(word) > 2:
                v[int(hashlib.md5(word.encode()).hexdigest(), 16) % self.dim] += 1
        n = np.linalg.norm(v)
        return v / n if n else v

    def embed_passages(self, texts):
        return np.array([self._vec(t) for t in texts])

    def embed_query(self, text):
        return self._vec(text)


def make_pdf(pages: list[str]) -> bytes:
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=A4)
    for text in pages:
        y = 800
        for line in text.split("\n"):
            c.drawString(50, y, line)
            y -= 16
        c.showPage()
    c.save()
    return buf.getvalue()


def make_docx(sections: list[tuple[str, list[str]]], table: list[list[str]] | None = None) -> bytes:
    doc = Document()
    for heading, paras in sections:
        doc.add_heading(heading, level=1)
        for p in paras:
            doc.add_paragraph(p)
    if table:
        t = doc.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for col, val in enumerate(row):
                t.cell(r, col).text = val
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def make_xlsx(sheets: dict[str, list[list]]) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    for title, rows in sheets.items():
        ws = wb.create_sheet(title)
        for row in rows:
            ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@pytest.fixture
def settings():
    return Settings(
        key_mode="user", server_api_key="", server_models=["claude-opus-5", "claude-haiku-4-5"],
        daily_question_limit=0,
        access_passcode=PASSCODE, token_secret="test-secret", access_token_ttl_minutes=60,
        rate_limit_auth_per_15min=100, allowed_providers=["anthropic", "openai", "gemini", "openai_compatible"],
        embedding_model="unused", allowed_origins=["http://localhost:8000"], max_file_mb=5, max_zip_mb=5, max_files=10,
        max_total_pages=100, session_ttl_minutes=60, rate_limit_ask_per_min=100,
        rate_limit_ingest_per_hour=100, llm_timeout_s=10,
    )


@pytest.fixture
def report_pdf():
    return make_pdf([
        "Quarterly report\nThis document covers company performance.",
        "Revenue\nThe Q2 revenue target was 48 crore rupees.\nMarketing spend rose slightly.",
        "Hiring\nThe team hired twelve engineers in Bengaluru.",
    ])


@pytest.fixture
def policy_docx():
    return make_docx(
        [("Leave policy", ["Employees get 24 days of paid leave each year."]),
         ("Travel policy", ["International travel needs approval from a director."])],
        table=[["Grade", "Hotel limit"], ["L1", "5000"], ["L2", "8000"]],
    )


@pytest.fixture
def sales_xlsx():
    import datetime as dt
    return make_xlsx({
        "Revenue": [["Region", "Q1 revenue", "Q2 revenue", "Updated"],
                    ["North", 10.5, 11.0, dt.date(2026, 7, 1)],
                    ["South", 12.4, 14.0, dt.date(2026, 7, 1)],
                    [None, None, None, None],
                    ["West", 9, 9.5, dt.date(2026, 7, 2)]],
        "Notes": [["Topic", "Comment"], ["Pricing", "Discounts capped at 12 percent"]],
    })
