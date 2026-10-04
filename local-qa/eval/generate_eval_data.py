"""Generates a dummy evaluation pack for Local Q&A, plus the answer key, deterministically.

Use case A (unstructured -> insights and tables):
  data/A_Unstructured/Customer_Calls_Sep2026.docx   12 customer call notes
  data/A_Unstructured/Store_Visit_Reports_Aug2026.pdf  6 store visit reports, one per page
Use case B (product-level sales -> summaries, charts):
  data/B_Sales/Sales_FY27_H1.xlsx        184 monthly product x region rows + product targets
  data/B_Sales/Sales_Review_H1_Commentary.docx  management commentary
Plus Evaluation_Scoresheet.xlsx: questions, reference answers computed from the same data,
scoring rubric, and blank columns to fill while testing.

Run:  local-qa/backend/.venv/bin/python local-qa/eval/generate_eval_data.py
All names, companies and figures are fictitious.
"""
import random
from collections import defaultdict
from pathlib import Path

from docx import Document
from docx.shared import Pt
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas

OUT = Path(__file__).resolve().parent
A_DIR = OUT / "data" / "A_Unstructured"
B_DIR = OUT / "data" / "B_Sales"
A_DIR.mkdir(parents=True, exist_ok=True)
B_DIR.mkdir(parents=True, exist_ok=True)
rng = random.Random(2026)

# ============================================================== use case B: sales data

PRODUCTS = [  # sku, product, category, unit price (INR), base monthly units per region
    ("AUR-01", "Aurora Kettle", "Kitchen", 2499, 420),
    ("BRZ-02", "Breeze Fan", "Home comfort", 3199, 380),
    ("CAS-03", "Cascade Water Purifier", "Kitchen", 12999, 140),
    ("DRF-04", "Drift Air Cooler", "Home comfort", 8499, 160),
    ("EMB-05", "Ember Room Heater", "Home comfort", 2899, 90),
    ("FLX-06", "Flux Mixer Grinder", "Kitchen", 4599, 210),
    ("GLW-07", "Glow LED Lamp", "Lighting", 899, 650),
    ("HAL-08", "Halo Smart Bulb", "Lighting", 1299, 0),
]
MONTHS = ["2026-04", "2026-05", "2026-06", "2026-07", "2026-08", "2026-09"]
MONTH_NAMES = {"2026-04": "April", "2026-05": "May", "2026-06": "June", "2026-07": "July", "2026-08": "August", "2026-09": "September"}
REGIONS = {"North": 1.10, "South": 1.00, "East": 0.80, "West": 1.15}
SEASON = {
    "Breeze Fan": [1.3, 1.5, 1.4, 1.0, 0.8, 0.7],
    "Drift Air Cooler": [1.6, 2.2, 1.8, 0.9, 0.6, 0.5],  # May heatwave
    "Ember Room Heater": [0.5, 0.3, 0.2, 0.3, 0.6, 1.2],
}
HALO_LAUNCH = [0, 0, 120, 260, 380, 450]  # launched June 2026
DEFAULT_DISCOUNT = {"Kitchen": 0.05, "Home comfort": 0.08, "Lighting": 0.10}

rows = []
for m_i, month in enumerate(MONTHS):
    for region, rf in REGIONS.items():
        for sku, name, cat, price, base in PRODUCTS:
            if name == "Halo Smart Bulb":
                if HALO_LAUNCH[m_i] == 0:
                    continue  # not launched yet
                units = HALO_LAUNCH[m_i] * rf
            else:
                units = base * rf * SEASON.get(name, [1] * 6)[m_i]
            if name == "Cascade Water Purifier" and region == "South":
                units *= {"2026-08": 0.25, "2026-09": 1.10}.get(month, 1)  # August stock-out
            units = max(1, round(units * rng.uniform(0.92, 1.08)))
            disc = DEFAULT_DISCOUNT[cat]
            if name == "Ember Room Heater" and month in ("2026-04", "2026-05"):
                disc = 0.25  # clearance
            if name == "Halo Smart Bulb" and month == "2026-06":
                disc = 0.15  # launch offer
            net = round(units * price * (1 - disc))
            rows.append({"Month": month, "Region": region, "SKU": sku, "Product": name, "Category": cat,
                         "Units": units, "Unit price (INR)": price, "Discount": disc, "Net revenue (INR)": net})

def total(key, value_col, filt=lambda r: True):
    out = defaultdict(float)
    for r in rows:
        if filt(r):
            out[r[key]] += r[value_col]
    return dict(out)

units_by_product = total("Product", "Units")
rev_by_product = total("Product", "Net revenue (INR)")
rev_by_region = total("Region", "Net revenue (INR)")
rev_by_category = total("Category", "Net revenue (INR)")
TARGET_FACTOR = {"Aurora Kettle": 0.95, "Breeze Fan": 0.90, "Cascade Water Purifier": 1.15, "Drift Air Cooler": 0.85,
                 "Ember Room Heater": 1.20, "Flux Mixer Grinder": 1.05, "Glow LED Lamp": 0.97, "Halo Smart Bulb": 1.30}
targets = {p: int(round(units_by_product[p] * TARGET_FACTOR[p], -2)) for p in units_by_product}

wb = Workbook()
ws = wb.active
ws.title = "Transactions"
headers = list(rows[0].keys())
ws.append(headers)
for r in rows:
    ws.append([r[h] for h in headers])
pm = wb.create_sheet("Products")
pm.append(["SKU", "Product", "Category", "Unit price (INR)", "H1 unit target", "Launch month"])
for sku, name, cat, price, _ in PRODUCTS:
    pm.append([sku, name, cat, price, targets[name], "2026-06" if name == "Halo Smart Bulb" else "Before 2026-04"])
for sheet in (ws, pm):
    for c in sheet[1]:
        c.font = Font(name="Arial", bold=True)
    for row in sheet.iter_rows(min_row=2):
        for c in row:
            c.font = Font(name="Arial")
    for i, h in enumerate([c.value for c in sheet[1]], start=1):
        sheet.column_dimensions[get_column_letter(i)].width = max(12, len(str(h)) + 4)
for c in ws["H"][1:]:
    c.number_format = "0%"
for col in ("G", "I"):
    for c in ws[col][1:]:
        c.number_format = "#,##0"
wb.save(B_DIR / "Sales_FY27_H1.xlsx")

# Commentary (qualitative; consistent with the data)
doc = Document()
doc.styles["Normal"].font.name = "Arial"
doc.styles["Normal"].font.size = Pt(10.5)
doc.add_heading("Sales Review H1 FY27 (April to September 2026)", 0)
for heading, paras in [
    ("Summary", ["H1 was shaped by an unusually strong summer and one supply problem. Home comfort products led the "
                 "first quarter, kitchen appliances held steady, and the new Halo Smart Bulb ramped up after its June launch."]),
    ("Summer demand", ["A heatwave across North and West India in May pushed Drift Air Cooler sales to their highest level of "
                       "the half. Breeze Fan followed the same pattern. Both products eased from July as the monsoon set in."]),
    ("Supply issue: Cascade Water Purifier", ["Our purifier supplier missed two shipments in late July. Southern stores ran out "
                                              "of the Cascade Water Purifier for most of August, and sales in the South fell sharply "
                                              "that month before recovering in September once stock arrived."]),
    ("Halo Smart Bulb launch", ["Halo Smart Bulb launched in June with a 15% introductory discount. Uptake grew every month, but "
                                "customers reported difficulties pairing the bulb with the mobile app, which the product team is fixing."]),
    ("Pricing and discounts", ["Ember Room Heater was cleared at 25% off in April and May to reduce winter stock. Standard "
                               "discounts were 5% for kitchen, 8% for home comfort and 10% for lighting products."]),
    ("Outlook for H2", ["We expect heater demand to rise from October. Priorities are restoring purifier supply in the South, "
                        "fixing Halo app pairing, and staff training on purifier demonstrations in the East."]),
]:
    doc.add_heading(heading, level=1)
    for p in paras:
        doc.add_paragraph(p)
doc.save(B_DIR / "Sales_Review_H1_Commentary.docx")

# ============================================================== use case A: unstructured notes

CALLS = [  # customer, city, date, product, issue (category), sentiment, owner, due, order, churn
    ("Meridian Foods", "Pune", "2 Sep 2026", "Cascade Water Purifier", "Filter cartridge availability", "Negative", "Ravi Menon", "10 Sep 2026", "40 units on hold", True,
     "Spoke with Anita Desai, procurement head at Meridian Foods in Pune, on 2 September. She said replacement filter cartridges for the "
     "Cascade Water Purifier have been unavailable for three weeks and their canteen units are running on expired filters. She is unhappy "
     "and said they are 'evaluating AquaPure' if supply does not improve. A planned order of 40 purifiers is on hold. Ravi Menon will "
     "confirm a cartridge delivery date by 10 September."),
    ("Sunrise Retail", "Mumbai", "3 Sep 2026", "Halo Smart Bulb", "App pairing problems", "Neutral", "Priya Shah", "15 Sep 2026", "Reorder of 300 units likely", False,
     "Call with Karan Mehta, category manager at Sunrise Retail, Mumbai, on 3 September. Halo Smart Bulb sells well in their stores, but "
     "about one in ten buyers come back because the bulb will not pair with the app. He is otherwise positive and expects to reorder "
     "around 300 units if the fix ships soon. Priya Shah to share the app update timeline by 15 September."),
    ("GreenLeaf Hotels", "Bengaluru", "4 Sep 2026", "Drift Air Cooler", "Delivery delays", "Negative", "Arjun Rao", "8 Sep 2026", "None", False,
     "GreenLeaf Hotels (Bengaluru) facilities manager Suresh Kumar called on 4 September about 25 Drift Air Coolers ordered in July that "
     "arrived eleven days late. He was frustrated because the summer peak had passed. No new order this quarter. Arjun Rao will send a "
     "written apology and a credit note proposal by 8 September."),
    ("Kaveri Distributors", "Chennai", "5 Sep 2026", "Cascade Water Purifier", "Filter cartridge availability", "Negative", "Ravi Menon", "12 Sep 2026", "None", True,
     "Kaveri Distributors, Chennai, spoke to us on 5 September. Owner Lakshmi Iyer said retailers keep asking for Cascade filter "
     "cartridges and for purifier stock, which ran out in August. She mentioned that AquaPure is offering better margins and she may "
     "move part of her business to them. Ravi Menon to call back with a stock plan by 12 September."),
    ("Urban Nest Interiors", "Delhi", "8 Sep 2026", "Glow LED Lamp", "None (positive feedback)", "Positive", "Neha Gupta", "30 Sep 2026", "500 units", False,
     "Positive call with Urban Nest Interiors in Delhi on 8 September. Designer Rahul Khanna likes the Glow LED Lamp finish and wants "
     "500 units for a housing project in October. Neha Gupta to send a project quote by 30 September."),
    ("Coastal Mart", "Kochi", "9 Sep 2026", "Breeze Fan", "Price higher than competitor", "Neutral", "Arjun Rao", "20 Sep 2026", "150 units if price matched", False,
     "Coastal Mart, Kochi, purchase lead Thomas Mathew said on 9 September that CoolWind fans are about 10% cheaper than the Breeze Fan. "
     "He would order 150 Breeze Fans if we can come closer on price. Arjun Rao to check a volume discount by 20 September."),
    ("Sapphire Electronics", "Hyderabad", "11 Sep 2026", "Flux Mixer Grinder", "Delivery delays", "Negative", "Priya Shah", "18 Sep 2026", "None", False,
     "Sapphire Electronics in Hyderabad (owner Imran Baig) complained on 11 September that their August order of Flux Mixer Grinders "
     "came a week late and two cartons were damaged. Priya Shah will arrange replacements and confirm by 18 September."),
    ("Himalaya Stores", "Dehradun", "12 Sep 2026", "Ember Room Heater", "None (pre-season order)", "Positive", "Neha Gupta", "25 Sep 2026", "220 units for October", False,
     "Himalaya Stores, Dehradun, buyer Pooja Negi called on 12 September to book 220 Ember Room Heaters for early October delivery "
     "ahead of winter. Happy with last year's heaters. Neha Gupta to confirm the delivery slot by 25 September."),
    ("TechHome Solutions", "Bengaluru", "15 Sep 2026", "Halo Smart Bulb", "App pairing problems", "Negative", "Priya Shah", "22 Sep 2026", "Paused", True,
     "TechHome Solutions, Bengaluru, installs smart-home kits. Founder Vikram Joshi said on 15 September that Halo pairing failures "
     "are causing installation call-backs and he is trialling a competitor bulb from LumaSmart. Orders are paused. Priya Shah to "
     "arrange an on-site session with the app team by 22 September."),
    ("Annapurna Caterers", "Pune", "16 Sep 2026", "Aurora Kettle", "Delivery delays", "Neutral", "Arjun Rao", "19 Sep 2026", "60 units", False,
     "Annapurna Caterers, Pune, kitchen manager Meera Kulkarni said on 16 September that their 60 Aurora Kettles were delivered four "
     "days late but work well. She will reorder in November. Arjun Rao to share the revised dispatch schedule by 19 September."),
    ("BrightPath Schools", "Lucknow", "18 Sep 2026", "Halo Smart Bulb", "App pairing problems", "Neutral", "Priya Shah", "29 Sep 2026", "Pilot of 80 units", False,
     "BrightPath Schools in Lucknow (admin officer Sanjay Verma), 18 September: their pilot of 80 Halo Smart Bulbs had a few pairing "
     "issues, but staff found a workaround. They will decide on a full rollout after the app update. Priya Shah to follow up by 29 September."),
    ("Delta Supermarkets", "Ahmedabad", "22 Sep 2026", "Cascade Water Purifier", "Filter cartridge availability; Price higher than competitor", "Neutral", "Ravi Menon", "1 Oct 2026", "25 units", False,
     "Delta Supermarkets, Ahmedabad, category head Nikhil Patel said on 22 September that shoppers ask for spare filters they cannot "
     "find, and that AquaPure purifiers are priced about ₹1,500 lower. He still plans to order 25 Cascade purifiers. Ravi Menon to "
     "send a filter supply commitment by 1 October."),
]
doc = Document()
doc.styles["Normal"].font.name = "Arial"
doc.styles["Normal"].font.size = Pt(10.5)
doc.add_heading("Customer call notes: September 2026", 0)
doc.add_paragraph("Notes typed by the key accounts team after each call. Informal; not reviewed.")
for i, c in enumerate(CALLS, start=1):
    doc.add_heading(f"Call {i}: {c[0]}, {c[1]} ({c[2]})", level=2)
    doc.add_paragraph(c[10])
doc.save(A_DIR / "Customer_Calls_Sep2026.docx")

STORES = [  # store, region, date, stock-outs [(product, days)], competitor note, report text
    ("Chennai, Anna Nagar", "South", "12 Aug 2026", [("Cascade Water Purifier", 9)], "AquaPure purifier promoted at ₹11,499",
     "Visit by area manager on 12 August 2026. The store has been out of the Cascade Water Purifier for 9 days; staff are sending "
     "customers to other branches. Displays are clean and well lit. A nearby dealer is promoting the AquaPure purifier at ₹11,499, "
     "about ₹1,500 below Cascade. Footfall estimated at 180 visitors a day. Action: escalate purifier stock to the regional warehouse."),
    ("Bengaluru, Indiranagar", "South", "14 Aug 2026", [("Cascade Water Purifier", 6)], None,
     "Visit on 14 August 2026. Cascade Water Purifier out of stock for 6 days. The Halo Smart Bulb display has no working demo phone, "
     "so staff cannot show the app; several shoppers asked about pairing. Footfall about 220 a day. Actions: send a demo phone; "
     "restock purifiers."),
    ("Delhi, Saket", "North", "18 Aug 2026", [], "CoolWind fans cut prices by about 10%",
     "Visit on 18 August 2026. No stock-outs. Drift Air Coolers are overstocked now that summer demand has faded; recommend moving 40 "
     "units to the outlet channel. Ember Room Heaters have started arriving early for winter. CoolWind has cut fan prices by about 10%. "
     "Footfall about 260 a day."),
    ("Mumbai, Andheri", "West", "20 Aug 2026", [("Glow LED Lamp", 4)], None,
     "Visit on 20 August 2026. Strong interest in the Halo Smart Bulb, but staff report frequent pairing complaints from buyers. Glow "
     "LED Lamp has been out of stock for 4 days after a bulk corporate order. Footfall about 300 a day, the highest of the stores visited."),
    ("Kolkata, Salt Lake", "East", "22 Aug 2026", [], None,
     "Visit on 22 August 2026. Footfall low, about 90 a day, because of heavy rain. No stock-outs. New staff are not confident "
     "demonstrating the Cascade Water Purifier; recommend a training session. Displays need fresh price labels."),
    ("Hyderabad, Banjara Hills", "South", "25 Aug 2026", [("Cascade Water Purifier", 11)], "AquaPure running a festive promotion",
     "Visit on 25 August 2026. Cascade Water Purifier out of stock for 11 days, the third southern store with this problem. AquaPure "
     "is running a festive promotion in the same mall. Flux Mixer Grinder display carton damaged in transit. Footfall about 200 a day. "
     "Action: priority purifier allocation for South stores."),
]
pdfmetrics.registerFont(TTFont("DejaVu", "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"))
pdfmetrics.registerFont(TTFont("DejaVu-Bold", "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"))
pdf = canvas.Canvas(str(A_DIR / "Store_Visit_Reports_Aug2026.pdf"), pagesize=A4)
pdf.setTitle("Store visit reports, August 2026")
for store, region, date, _, _, text in STORES:
    y = 270 * mm
    pdf.setFont("DejaVu-Bold", 15)
    pdf.drawString(20 * mm, y, f"Store visit report: {store}")
    pdf.setFont("DejaVu", 10.5)
    y -= 9 * mm
    pdf.drawString(20 * mm, y, f"Region: {region}    Visit date: {date}")
    y -= 10 * mm
    words, line = text.split(), ""
    for w in words:  # simple wrap at ~95 characters
        if len(line) + len(w) + 1 > 95:
            pdf.drawString(20 * mm, y, line)
            y -= 6 * mm
            line = w
        else:
            line = f"{line} {w}".strip()
    pdf.drawString(20 * mm, y, line)
    pdf.showPage()
pdf.save()

# ============================================================== answer key + scoresheet

inr = lambda v: f"₹{v:,.0f}"
lakh = lambda v: f"₹{v / 1e5:,.1f} lakh"
top_product = max(rev_by_product, key=rev_by_product.get)
missed = sorted([p for p in targets if units_by_product[p] < targets[p]])
cooler_monthly = {MONTH_NAMES[m]: int(sum(r["Units"] for r in rows if r["Product"] == "Drift Air Cooler" and r["Month"] == m)) for m in MONTHS}
cas_south = {MONTH_NAMES[m]: int(sum(r["Units"] for r in rows if r["Product"] == "Cascade Water Purifier" and r["Region"] == "South" and r["Month"] == m)) for m in MONTHS}
halo_monthly = {MONTH_NAMES[m]: int(sum(r["Units"] for r in rows if r["Product"] == "Halo Smart Bulb" and r["Month"] == m)) for m in MONTHS[2:]}
grand_total = sum(r["Net revenue (INR)"] for r in rows)
fmt_map = lambda d, f=lambda v: f"{int(v):,}": "; ".join(f"{k} {f(v)}" for k, v in d.items())

call_table = "\n".join(f"{c[0]} | {c[1]} | {c[2]} | {c[3]} | {c[4]} | {c[5]} | {c[6]} | {c[7]}" for c in CALLS)
stockouts = "\n".join(f"{s[0]} | {s[1]} | {p} | {d} days | {s[2]}" for s in STORES for p, d in s[3])

CASES = [
    # id, use case, type, question, reference answer, key facts, pass criteria, settings
    ("A1", "A: Unstructured", "Extraction to table",
     "Create a table of all customer calls with customer, city, date, product, main issue, sentiment, follow-up owner and due date.",
     "A 12-row table:\nCustomer | City | Date | Product | Main issue | Sentiment | Owner | Due\n" + call_table,
     "All 12 customers present; owners and due dates match; sentiment reasonable (Negative for Meridian, GreenLeaf, Kaveri, Sapphire, TechHome).",
     "2 = all 12 rows correct; 1 = 9-11 rows or minor field errors; 0 = <9 rows or invented rows.",
     "Default settings (Answer from: Auto reads all documents); Max answer length 1500+"),
    ("A2", "A: Unstructured", "Insight",
     "What are the most common complaints across the customer calls? Count how many calls mention each.",
     "Halo app pairing problems: 3 calls (Sunrise, TechHome, BrightPath). Cascade filter cartridge availability: 3 (Meridian, Kaveri, Delta). "
     "Delivery delays: 3 (GreenLeaf, Sapphire, Annapurna). Price higher than competitor: 2 (Coastal Mart vs CoolWind, Delta vs AquaPure).",
     "Three themes tied at 3 calls each; price as a 4th theme with 2.", "2 = all themes with right counts; 1 = right themes, wrong counts; 0 = wrong themes.",
     "Default settings (Answer from: Auto)"),
    ("A3", "A: Unstructured", "Insight",
     "Which customers are at risk of moving to a competitor, and why?",
     "Meridian Foods (Pune): no filter cartridges, evaluating AquaPure, 40-unit order on hold. Kaveri Distributors (Chennai): stock and "
     "cartridge shortage, AquaPure offering better margins. TechHome Solutions (Bengaluru): Halo pairing failures, trialling LumaSmart, orders paused.",
     "Exactly these 3; competitor names AquaPure, AquaPure, LumaSmart.", "2 = all 3 with reasons; 1 = 2 of 3; 0 = fewer or invented.",
     "Default settings (Answer from: Auto)"),
    ("A4", "A: Unstructured", "Extraction to table",
     "List the follow-up actions due on or before 15 September 2026, with owner and due date, as a table.",
     "GreenLeaf Hotels: apology + credit note, Arjun Rao, 8 Sep. Meridian Foods: cartridge delivery date, Ravi Menon, 10 Sep. "
     "Kaveri Distributors: stock plan, Ravi Menon, 12 Sep. Sunrise Retail: app update timeline, Priya Shah, 15 Sep.",
     "Exactly 4 actions.", "2 = exactly these 4; 1 = 3 correct or 1 extra; 0 otherwise.", "Default settings (Answer from: Auto)"),
    ("A5", "A: Unstructured", "Extraction to table",
     "From the store visit reports, make a table of stock-outs: store, region, product, how many days, visit date.",
     "Store | Region | Product | Days | Visit date\n" + stockouts,
     "4 rows: 3 Cascade Water Purifier stock-outs in South stores + Glow LED Lamp in Mumbai.", "2 = all 4 correct; 1 = 3 correct; 0 otherwise.",
     "Default settings"),
    ("A6", "A: Unstructured", "Summary",
     "Summarise what competitors are doing according to the store visits and customer calls.",
     "AquaPure: purifier at ₹11,499 (about ₹1,500 below Cascade) in Chennai, festive promotion in Hyderabad, better margins offered to "
     "Kaveri, mentioned by Meridian and Delta. CoolWind: fans about 10% cheaper (Delhi store, Coastal Mart). LumaSmart: smart bulb trial at TechHome.",
     "AquaPure, CoolWind, LumaSmart with the right products.", "2 = all three; 1 = two; 0 = one or invented competitors.",
     "Default settings (Answer from: Auto)"),
    ("A7", "Cross-use-case", "Cross-document insight",
     "Do the store visit reports or customer calls explain any dip in the sales data?",
     f"Yes: Cascade Water Purifier sales in the South fell in August ({cas_south['July']:,} units in July to {cas_south['August']:,} in August) "
     "because South stores (Chennai, Bengaluru, Hyderabad) were out of stock; the commentary blames missed supplier shipments. Sales recovered "
     f"in September ({cas_south['September']:,}). Customers Kaveri and Meridian also cite cartridge/stock shortages.",
     "Links the South August purifier dip to the stock-outs, cites both sales data and visit reports.",
     "2 = link made with evidence from both sources; 1 = link made, one source; 0 = no link or wrong cause.",
     "Load A and B folders together; default settings"),
    ("A8", "A: Unstructured", "Negative (not in documents)",
     "What did customers say about the warranty extension programme?",
     "Not in the documents. The tool should answer that it couldn't find this.",
     "Must not invent anything.", "2 = says not found; 0 = any invented content.", "Default settings"),
    ("A9", "A: Unstructured", "Derived count",
     "How many store visits were there per month?",
     "6 store visits, all in August 2026 (12, 14, 18, 20, 22 and 25 August); none in any other month.",
     "6 visits, all August.", "2 = 6 in August and none elsewhere; 1 = August only or wrong count; 0 = 'not found' or invented months.",
     "Default settings (found in real testing: an over-strict prompt made the tool answer 'not found')"),
    ("B1", "B: Sales", "Lookup",
     "What discount was given on the Ember Room Heater in April 2026, and why?",
     "25% clearance discount (April and May) to reduce winter stock; standard home comfort discount is 8%.",
     "25%, clearance reason.", "2 = both; 1 = number only; 0 = wrong.", "Default settings"),
    ("B2", "B: Sales", "Narrative",
     "Why did Drift Air Cooler sales peak in May?",
     "A heatwave across North and West India in May (commentary); sales eased from July with the monsoon.",
     "Heatwave, May.", "2 = cause + timing; 1 = cause only; 0 = wrong.", "Default settings"),
    ("B3", "B: Sales", "Chart (exact numbers)",
     "Plot monthly units sold of Drift Air Cooler as a line chart.",
     "Proposal table, then line chart with all regions summed: " + fmt_map(cooler_monthly) + ".",
     "Values match exactly; source shows Sales_FY27_H1.xlsx sheet Transactions.", "2 = exact values and plotted after confirm; 1 = right shape, wrong grouping; 0 = wrong numbers.",
     "Default settings"),
    ("B4", "B: Sales", "Chart (exact numbers)",
     "Bar chart of total net revenue by product for H1.",
     "Bar chart, values (₹): " + fmt_map(dict(sorted(rev_by_product.items(), key=lambda kv: -kv[1])), inr) + f". Top: {top_product}.",
     "Exact values (page-computed).", "2 = exact; 0 = otherwise.", "Default settings"),
    ("B5", "B: Sales", "Chart (exact numbers)",
     "Show net revenue by region as a bar chart.",
     "Values (₹): " + fmt_map(dict(sorted(rev_by_region.items(), key=lambda kv: -kv[1])), inr) + ".",
     "Exact values.", "2 = exact; 0 = otherwise.", "Default settings"),
    ("B6", "B: Sales", "Aggregate (text answer)",
     "Which product earned the most net revenue in H1, and how much?",
     f"{top_product}, {inr(rev_by_product[top_product])} ({lakh(rev_by_product[top_product])}).",
     "Product name and amount within 1%.",
     "2 = right product and amount; 1 = right product, wrong/missing amount, or says it can only see part of the data; 0 = wrong product stated confidently.",
     "Default settings; check the 'Exact figures calculated by the page' table under the answer"),
    ("B7", "B: Sales", "Aggregate (text answer)",
     "What was total net revenue for H1 across all products and regions?",
     f"{inr(grand_total)} ({lakh(grand_total)}).",
     "Within 1%, or an honest statement that the full total can't be computed from the passages it sees.",
     "2 = correct; 1 = honest 'can't compute all rows'; 0 = a wrong number stated as fact.", "Default settings; check the calculated table under the answer"),
    ("B8", "B: Sales", "Aggregate (text answer)",
     "Which products missed their H1 unit target?",
     "Missed: " + ", ".join(f"{p} ({int(units_by_product[p]):,} vs target {targets[p]:,})" for p in missed) + ".",
     "Exactly these products.", "2 = exact list; 1 = partial or honest limitation; 0 = wrong list stated as fact.",
     "Default settings; check the calculated table under the answer"),
    ("B9", "B: Sales", "Summary",
     "Summarise how the Halo Smart Bulb has done since launch.",
     "Launched June 2026 with a 15% introductory discount; units grew every month (" + fmt_map(halo_monthly) + "); app pairing problems reported "
     "by customers and stores; app fix in progress.", "Launch month, discount, growth, pairing issue.",
     "2 = all four; 1 = two or three; 0 = otherwise.", "Default settings"),
    ("B10", "B: Sales", "Negative (not in documents)",
     "What were sales in October 2026?",
     "Not in the data (the data covers April to September 2026). Should say it couldn't find this.",
     "Must not invent October figures.", "2 = says not found / data ends September; 0 = any invented figure.", "Default settings"),
]

wb = Workbook()
ARIAL = "Arial"
bold = Font(name=ARIAL, bold=True)
normal = Font(name=ARIAL)
yellow = PatternFill("solid", fgColor="FFFF00")
grey = PatternFill("solid", fgColor="EEF0F2")
thin = Side(style="thin", color="D1D5DB")
box = Border(left=thin, right=thin, top=thin, bottom=thin)
wrap = Alignment(wrap_text=True, vertical="top")

ins = wb.active
ins.title = "Instructions"
lines = [
    ("Local Q&A evaluation scoresheet", True),
    ("", False),
    ("How to run", True),
    ("1. Open https://ng-04.github.io/news_sentiment/#local-qa and enter your Claude API key (model: claude-opus-5).", False),
    ("2. Use case A: add the folder data/A_Unstructured. Use case B: add data/B_Sales. For A7, add both folders together.", False),
    ("3. Before each question, set the Advanced settings shown in the 'Settings to use' column (re-index if chunk size changes).", False),
    ("4. Ask each question exactly as written, in a fresh chat (Clear chat) unless it is a follow-up.", False),
    ("5. Paste the answer into 'Actual answer', time it, and score it with the rubric. For chart questions, compare the proposal table with the reference before clicking Plot chart.", False),
    ("", False),
    ("Which cells to fill in", True),
    ("Yellow cells in the 'Test cases' sheet: Actual answer, Score (0-2), Citations correct (Y/N), Made-up facts (Y/N), Seconds, Notes.", False),
    ("Row 2 is a filled-in EXAMPLE showing the expected format; it is excluded from the Summary.", False),
    ("", False),
    ("Scoring rubric", True),
    ("Score: 2 = correct and complete; 1 = partly correct, or honestly states a limitation; 0 = wrong, incomplete, or invents facts.", False),
    ("Citations correct: Y if the cited file/page/sheet actually supports the answer.", False),
    ("Made-up facts: Y if the answer states anything that is not in the documents (a hallucination). Any Y here is a serious failure.", False),
    ("", False),
    ("About the reference answers", True),
    ("Reference answers were computed by generate_eval_data.py from the same dummy data (seed 2026), so they are exact. All names and figures are fictitious.", False),
    ("B6-B8 need totals across all 184 sales rows. The tool asks the page to calculate them exactly (shown as 'Exact figures calculated by the page' "
     "under the answer); the answer should quote those figures. Charts (B3-B5) use the same exact calculations.", False),
    ("With small document sets, 'Answer from: Auto' sends Claude the complete documents, so 'list all' and 'table of every' questions (A1-A4) can be complete.", False),
]
for i, (text, is_head) in enumerate(lines, start=1):
    c = ins.cell(row=i, column=1, value=text)
    c.font = Font(name=ARIAL, bold=is_head, size=14 if i == 1 else 10)
    c.alignment = Alignment(wrap_text=True, vertical="top")
ins.column_dimensions["A"].width = 120

tc = wb.create_sheet("Test cases")
cols = ["ID", "Use case", "Question type", "Question", "Reference answer", "Key facts to check", "Pass criteria",
        "Settings to use", "Actual answer", "Score (0-2)", "Citations correct (Y/N)", "Made-up facts (Y/N)", "Seconds", "Notes"]
widths = [6, 16, 18, 42, 60, 36, 40, 30, 50, 10, 12, 12, 9, 30]
tc.append(cols)
for i, w in enumerate(widths, start=1):
    tc.column_dimensions[get_column_letter(i)].width = w
for c in tc[1]:
    c.font = bold
    c.fill = grey
    c.border = box
    c.alignment = wrap
example = ["EX", "Example", "Lookup", "What discount did Ember get in April?", "25% clearance", "25%", "2 = number + reason",
           "Default", "Ember Room Heater had a 25% clearance discount in April [1].", 2, "Y", "N", 6, "EXAMPLE row: excluded from the Summary"]
tc.append(example)
for c in tc[2]:
    c.font = Font(name=ARIAL, italic=True, color="6B7280")
    c.alignment = wrap
    c.border = box
for case in CASES:
    tc.append(list(case) + [None, None, None, None, None, None])
first, last = 3, 2 + len(CASES)
for row in tc.iter_rows(min_row=first, max_row=last):
    for j, c in enumerate(row, start=1):
        c.font = normal
        c.alignment = wrap
        c.border = box
        if j >= 9:
            c.fill = yellow
score_dv = DataValidation(type="whole", operator="between", formula1="0", formula2="2", allow_blank=True)
yn_dv = DataValidation(type="list", formula1='"Y,N"', allow_blank=True)
tc.add_data_validation(score_dv)
tc.add_data_validation(yn_dv)
score_dv.add(f"J{first}:J{last}")
yn_dv.add(f"K{first}:L{last}")
tc.freeze_panes = "E2"
tc["E1"].comment = Comment("Computed from the dummy data by generate_eval_data.py (seed 2026).", "Local Q&A")

sm = wb.create_sheet("Summary")
rng_id = f"'Test cases'!$A${first}:$A${last}"
rng_uc = f"'Test cases'!$B${first}:$B${last}"
rng_sc = f"'Test cases'!$J${first}:$J${last}"
rng_ci = f"'Test cases'!$K${first}:$K${last}"
rng_mf = f"'Test cases'!$L${first}:$L${last}"
rng_s = f"'Test cases'!$M${first}:$M${last}"
sm.append(["Measure", "A: Unstructured", "B: Sales", "Cross-use-case", "All", "Target"])
groups = ["A: Unstructured", "B: Sales", "Cross-use-case"]
def per_group(tmpl):
    return [tmpl.format(g=f'"{g}"') for g in groups]
rows_def = [
    ("Questions in the test set", per_group(f"=COUNTIF({rng_uc},{{g}})"), f"=COUNTA({rng_id})", ""),
    ("Questions scored", per_group(f"=COUNTIFS({rng_uc},{{g}},{rng_sc},\"<>\")"), f"=COUNT({rng_sc})", ""),
    ("Fully correct (score 2)", per_group(f"=COUNTIFS({rng_uc},{{g}},{rng_sc},2)"), f"=COUNTIF({rng_sc},2)", ""),
    ("Accuracy (points / max)", per_group(f"=IFERROR(SUMIFS({rng_sc},{rng_uc},{{g}})/(2*COUNTIFS({rng_uc},{{g}},{rng_sc},\"<>\")),\"\")"),
     f"=IFERROR(SUM({rng_sc})/(2*COUNT({rng_sc})),\"\")", "≥ 85%"),
    ("Citations correct", per_group(f"=IFERROR(COUNTIFS({rng_uc},{{g}},{rng_ci},\"Y\")/COUNTIFS({rng_uc},{{g}},{rng_ci},\"?*\"),\"\")"),
     f"=IFERROR(COUNTIF({rng_ci},\"Y\")/COUNTIF({rng_ci},\"?*\"),\"\")", "≥ 95%"),
    ("Answers with made-up facts", per_group(f"=COUNTIFS({rng_uc},{{g}},{rng_mf},\"Y\")"), f"=COUNTIF({rng_mf},\"Y\")", "0"),
    ("Average seconds per answer", per_group(f"=IFERROR(AVERAGEIFS({rng_s},{rng_uc},{{g}}),\"\")"), f"=IFERROR(AVERAGE({rng_s}),\"\")", "< 20"),
]
for label, cells, all_f, target in rows_def:
    sm.append([label, *cells, all_f, target])
for r in sm.iter_rows(min_row=1, max_row=sm.max_row):
    for c in r:
        c.font = bold if c.row == 1 or c.column == 1 else normal
        c.border = box
for c in sm[1]:
    c.fill = grey
for col in "BCDE":
    sm[f"{col}5"].number_format = "0%"
    sm[f"{col}6"].number_format = "0%"
    sm[f"{col}8"].number_format = "0.0"
sm.column_dimensions["A"].width = 32
for col in "BCDEF":
    sm.column_dimensions[col].width = 18
sm.cell(row=sm.max_row + 2, column=1, value="All values are formulas over the 'Test cases' sheet (the EXAMPLE row is outside the counted range).").font = Font(name=ARIAL, italic=True, color="6B7280")
wb.calculation.fullCalcOnLoad = True  # no cached values here; Excel calculates on open
wb.move_sheet("Summary", offset=-1)
wb.save(OUT / "Evaluation_Scoresheet.xlsx")

# ============================================================== website sample data
import json
import shutil

SAMPLES = OUT.parent / "samples"
FOLDERS = {"A_Unstructured": "Customer insights", "B_Sales": "Sales"}
DESCRIPTIONS = {
    "Customer_Calls_Sep2026.docx": ("Customer call notes", "12 informal notes from key-account calls: complaints, sentiment, competitor mentions, follow-ups."),
    "Store_Visit_Reports_Aug2026.pdf": ("Store visit reports", "6 one-page reports from store visits: stock-outs, displays, competitor pricing, footfall."),
    "Sales_FY27_H1.xlsx": ("Sales, April to September 2026", "184 monthly rows of units and net revenue by product and region, plus product targets."),
    "Sales_Review_H1_Commentary.docx": ("Sales commentary", "Management's written review of the half-year: heatwave, purifier supply problem, Halo launch."),
}
if SAMPLES.exists():
    shutil.rmtree(SAMPLES / "files", ignore_errors=True)
files_meta = []
for src_dir in (A_DIR, B_DIR):
    folder = FOLDERS[src_dir.name]
    for f in sorted(src_dir.iterdir()):
        dest = SAMPLES / "files" / folder / f.name
        dest.parent.mkdir(parents=True, exist_ok=True)
        shutil.copyfile(f, dest)
        title, desc = DESCRIPTIONS[f.name]
        files_meta.append({"name": f.name, "folder": folder, "path": f"files/{folder}/{f.name}", "title": title,
                           "description": desc, "kind": f.suffix.lstrip(".")})
GROUPS = {"A1": "Extract a table", "A5": "Extract a table", "A3": "Find insights", "A7": "Connect documents",
          "B6": "Calculate", "B8": "Calculate", "B3": "Chart", "B4": "Chart", "B9": "Summarise", "A8": "Trap question"}
by_id = {c[0]: c for c in CASES}
manifest = {
    "company": "Brightline Home Appliances, a fictitious company",
    "note": "All names, companies and figures are made up for this demo.",
    "files": files_meta,
    "questions": [{"id": i, "group": GROUPS[i], "text": by_id[i][3], "reference": by_id[i][4]} for i in GROUPS],
}
(SAMPLES / "manifest.json").write_text(json.dumps(manifest, indent=2, ensure_ascii=False) + "\n")

print(f"rows={len(rows)}  grand total={grand_total:,}  top={top_product} {rev_by_product[top_product]:,.0f}")
print("missed targets:", missed)
print("cooler monthly:", cooler_monthly)
print("cascade south:", cas_south)
